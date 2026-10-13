#!/usr/bin/env python3
import os
import re
import sys
import docx
from docx.shared import Inches, Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=80, bottom=80, left=120, right=120):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def set_table_borders(table, color="CCCCCC", sz="4", val="single"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(f'''
        <w:tblBorders {nsdecls("w")}>
            <w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:left w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:right w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:insideV w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
        </w:tblBorders>
    ''')
    tblPr.append(borders)

def set_cell_no_border(cell):
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = parse_xml(f'''
        <w:tcBorders {nsdecls("w")}>
            <w:top w:val="none"/>
            <w:left w:val="none"/>
            <w:bottom w:val="none"/>
            <w:right w:val="none"/>
        </w:tcBorders>
    ''')
    tcPr.append(tcBorders)

def add_formatted_text(paragraph, text, base_font_size=12, base_bold=False, base_italic=False, font_name="Times New Roman"):
    pattern = re.compile(r'(\*\*\*.*?\*\*\*|\*\*.*?\*\*|\*.*?\*|`.*?`)')
    tokens = pattern.split(text)
    
    for token in tokens:
        if not token:
            continue
        is_bold = base_bold
        is_italic = base_italic
        content = token
        
        if token.startswith('***') and token.endswith('***') and len(token) >= 6:
            is_bold = True
            is_italic = True
            content = token[3:-3]
        elif token.startswith('**') and token.endswith('**') and len(token) >= 4:
            is_bold = True
            content = token[2:-2]
        elif token.startswith('*') and token.endswith('*') and len(token) >= 2:
            is_italic = True
            content = token[1:-1]
        elif token.startswith('`') and token.endswith('`') and len(token) >= 2:
            content = token[1:-1]
            
        run = paragraph.add_run(content)
        run.bold = is_bold
        run.italic = is_italic
        run.font.name = font_name
        run.font.size = Pt(base_font_size)
        run.font.color.rgb = RGBColor(0, 0, 0)
        
        rPr = run._r.get_or_add_rPr()
        rFonts = parse_xml(f'<w:rFonts {nsdecls("w")} w:ascii="{font_name}" w:hAnsi="{font_name}" w:cs="{font_name}"/>')
        rPr.append(rFonts)

def main():
    root_dir = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    target_dir = os.path.join(root_dir, "docs", "Transformasi Digital")
    md_path = os.path.join(target_dir, "PROPOSAL_FULL.md")
    docx_path = os.path.join(target_dir, "Proposal Transformasi Digital - Grossmart.docx")
    figures_dir = os.path.join(target_dir, "figures")

    if not os.path.exists(md_path):
        print(f"Error: {md_path} tidak ditemukan.", file=sys.stderr)
        sys.exit(1)

    with open(md_path, "r", encoding="utf-8") as f:
        md_content = f.read()

    doc = docx.Document()

    # 1. Page Margins & Setup
    for section in doc.sections:
        section.page_width = Cm(21.59)
        section.page_height = Cm(27.94)
        section.top_margin = Cm(2.54)
        section.bottom_margin = Cm(2.54)
        section.left_margin = Cm(2.54)
        section.right_margin = Cm(2.54)

    # 2. Configure Normal style
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Times New Roman'
    normal_style.font.size = Pt(12)
    normal_style.font.color.rgb = RGBColor(0, 0, 0)
    normal_style.paragraph_format.line_spacing = 1.5
    normal_style.paragraph_format.space_after = Pt(3)
    normal_style.paragraph_format.space_before = Pt(0)

    abstrak_idx = md_content.find("## ABSTRAK")
    cover_text = md_content[:abstrak_idx].strip()
    body_text = md_content[abstrak_idx:].strip()

    # 3. Render Cover Page
    # Parse cover
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.line_spacing = 1.5
    p_title.paragraph_format.space_after = Pt(24)
    p_title.paragraph_format.space_before = Pt(36)
    add_formatted_text(p_title, "GROSSMART: APLIKASI KASIR DAN PENGELOLAAN OPERASIONAL TOKO KELONTONG YANG RAMAH PENGGUNA LANJUT USIA", base_font_size=12, base_bold=True)

    # Cover Template Icon
    cover_icon_path = os.path.join(figures_dir, "cover_icon.png")
    tpl_path = os.path.join(target_dir, "Template proposal - Transformasi Digital.docx")
    if os.path.exists(tpl_path):
        try:
            tpl = docx.Document(tpl_path)
            for rel in tpl.part.rels.values():
                if "image" in rel.target_ref:
                    with open(cover_icon_path, "wb") as f_icon:
                        f_icon.write(rel.target_part.blob)
                    break
        except Exception as e:
            print(f"Peringatan: Gagal ekstrak icon template: {e}", file=sys.stderr)

    if not os.path.exists(cover_icon_path):
        cover_icon_path = os.path.join(figures_dir, "template_cover_icon.png")
    if os.path.exists(cover_icon_path):
        p_icon = doc.add_paragraph()
        p_icon.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_icon.paragraph_format.line_spacing = 1.5
        p_icon.paragraph_format.space_before = Pt(6)
        p_icon.paragraph_format.space_after = Pt(12)
        run_icon = p_icon.add_run()
        with open(cover_icon_path, "rb") as f_icon_stream:
            run_icon.add_picture(f_icon_stream, width=Cm(5.29), height=Cm(5.29))

    p_space = doc.add_paragraph()
    p_space.paragraph_format.line_spacing = 1.5
    p_space.paragraph_format.space_after = Pt(0)
    p_space.paragraph_format.space_before = Pt(0)
    p_oleh = doc.add_paragraph()
    p_oleh.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_oleh.paragraph_format.line_spacing = 1.5
    p_oleh.paragraph_format.space_after = Pt(12)
    add_formatted_text(p_oleh, "Disusun Oleh :", base_font_size=12, base_bold=True)

    # Cover Table (Nama | Nim)
    tbl_cover = doc.add_table(rows=1, cols=2)
    tbl_cover.alignment = WD_TABLE_ALIGNMENT.CENTER
    
    # Borderless cover table
    tblPr = tbl_cover._tbl.tblPr
    borders = parse_xml(f'''
        <w:tblBorders {nsdecls("w")}>
            <w:top w:val="none"/>
            <w:left w:val="none"/>
            <w:bottom w:val="none"/>
            <w:right w:val="none"/>
            <w:insideH w:val="none"/>
            <w:insideV w:val="none"/>
        </w:tblBorders>
    ''')
    tblPr.append(borders)

    c0 = tbl_cover.cell(0, 0)
    c1 = tbl_cover.cell(0, 1)
    c0.width = Cm(7.0)
    c1.width = Cm(5.0)

    names = [
        "[DUMMY_REPLACE: Muhammad Fakhri Al-Faruq]",
        "[DUMMY_REPLACE: Budi Prasetyo Utomo]",
        "[DUMMY_REPLACE: Siti Rahmawati Dewi]"
    ]
    nims = [
        "[DUMMY_REPLACE: 220535601234]",
        "[DUMMY_REPLACE: 220535602345]",
        "[DUMMY_REPLACE: 220535603456]"
    ]

    p0 = c0.paragraphs[0]
    p0.paragraph_format.line_spacing = 1.5
    p0.paragraph_format.space_after = Pt(0)
    for idx, n in enumerate(names):
        if idx > 0:
            p0 = c0.add_paragraph()
            p0.paragraph_format.line_spacing = 1.5
            p0.paragraph_format.space_after = Pt(0)
        add_formatted_text(p0, n, base_font_size=12)

    p1 = c1.paragraphs[0]
    p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p1.paragraph_format.line_spacing = 1.5
    p1.paragraph_format.space_after = Pt(0)
    for idx, num in enumerate(nims):
        if idx > 0:
            p1 = c1.add_paragraph()
            p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p1.paragraph_format.line_spacing = 1.5
            p1.paragraph_format.space_after = Pt(0)
        add_formatted_text(p1, num, base_font_size=12)

    p_space2 = doc.add_paragraph()
    p_space2.paragraph_format.line_spacing = 1.5
    p_space2.paragraph_format.space_after = Pt(0)
    p_space2.paragraph_format.space_before = Pt(0)
    p_univ = doc.add_paragraph()
    p_univ.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_univ.paragraph_format.line_spacing = 1.5
    p_univ.paragraph_format.space_after = Pt(0)
    add_formatted_text(p_univ, "Universitas Bina Sarana Informatika", base_font_size=12)

    p_year = doc.add_paragraph()
    p_year.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_year.paragraph_format.line_spacing = 1.5
    p_year.paragraph_format.space_after = Pt(12)
    add_formatted_text(p_year, "2026", base_font_size=12)

    # Page break after cover
    doc.add_page_break()

    # 4. Parse Body Blocks
    lines = body_text.splitlines()
    i = 0
    
    # Table column width presets (total usable width = 16.51 cm)
    table_col_widths = {
        6: [Cm(4.01), Cm(2.5), Cm(2.5), Cm(2.5), Cm(2.5), Cm(2.5)], # Tabel 1 (6 cols)
        4: [Cm(3.51), Cm(3.5), Cm(5.0), Cm(4.5)],                    # Tabel 2, 4, 5 (4 cols)
        3: [Cm(3.51), Cm(6.5), Cm(6.5)],                            # Tabel 3, 7 (3 cols)
    }

    # Lampiran captions map
    lampiran_captions = {
        "figures/architecture.png": "Lampiran 1: Arsitektur Sistem — Gambar 1. Diagram Arsitektur Konseptual Sistem Grossmart. Pemisahan lima lapisan: klien React 18, caching PWA Service Worker mandiri, peladen Laravel 11, basis data MySQL InnoDB, dan pembatasan hak akses RBAC.",
        "figures/user_flow.png": "Lampiran 2: Alur Pengguna Terpadu — Gambar 2. Diagram Alur Kerja Kasir dan Pengawasan Pemilik. Siklus operasional dari buka kasir hingga tutup kasir buta.",
        "figures/screen_login.png": "Lampiran 3: Antarmuka Login Cepat Kasir — Gambar 3. Antarmuka Masuk Log (Login) dengan Pembagian Peran. Form autentikasi dan pemilihan akses kasir atau pemilik.",
        "figures/screen_pos.png": "Lampiran 4: Antarmuka Kasir POS Ergonomis Lansia — Gambar 4. Antarmuka Kasir Penjualan (POS) Ramah Lansia. Katalog visual, nominal belanja, tombol uang pas, dan pecahan Rupiah.",
        "figures/screen_shift_close.png": "Lampiran 5: Formulir Tutup Shift Buta — Gambar 5. Antarmuka Tutup Kasir Buta (Blind Cash-Up). Kalkulator hitung fisik uang dengan saldo sistem tersembunyi.",
        "figures/screen_dashboard.png": "Lampiran 6: Dashboard Pemilik Toko — Gambar 6. Dasbor Pengawasan Operasional dan Omzet Pemilik. Ringkasan omzet, laba kotor, selisih kas shift, dan log aktivitas operasional.",
        "figures/screen_products.png": "Lampiran 7: Manajemen Produk, Repack & Konsinyasi — Gambar 7. Pengelolaan Katalog Produk dan Konversi Eceran. Harga grosir, konversi stok dus ke eceran, dan informasi produk."
    }

    data_table_count = 0

    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        # Check for Table
        if stripped.startswith("|") and ("|" in stripped[1:]):
            table_lines = [line]
            i += 1
            while i < len(lines) and lines[i].strip().startswith("|"):
                table_lines.append(lines[i])
                i += 1
            
            # Parse table rows
            raw_rows = [r.strip() for r in table_lines if not re.match(r'^\|[-:\s|]+\|$', r.strip())]
            if raw_rows:
                data_table_count += 1
                header_cols = [c.strip() for c in raw_rows[0].split('|')[1:-1]]
                num_cols = len(header_cols)
                num_rows = len(raw_rows)
                
                table = doc.add_table(rows=num_rows, cols=num_cols)
                table.alignment = WD_TABLE_ALIGNMENT.CENTER
                set_table_borders(table, color="B0B0B0", sz="4")

                # Column widths
                widths = table_col_widths.get(num_cols)
                if not widths:
                    w_each = Cm(16.51 / num_cols)
                    widths = [w_each] * num_cols

                # Special width adjustments for SUS table (cols=4: No, Pernyataan, Aspek, Target)
                if num_cols == 4 and "No" in header_cols[0]:
                    widths = [Cm(1.01), Cm(8.0), Cm(5.0), Cm(2.5)]
                # Special width for Aksesibilitas table (cols=4)
                elif num_cols == 4 and "Aksesibilitas" in header_cols[0]:
                    widths = [Cm(3.51), Cm(3.8), Cm(6.7), Cm(2.5)]

                for r_idx, row_str in enumerate(raw_rows):
                    cols = [c.strip() for c in row_str.split('|')[1:-1]]
                    row_cells = table.rows[r_idx].cells
                    for c_idx in range(min(num_cols, len(cols))):
                        cell = row_cells[c_idx]
                        cell.width = widths[c_idx]
                        set_cell_margins(cell, top=60, bottom=60, left=90, right=90)
                        
                        p = cell.paragraphs[0]
                        p.paragraph_format.line_spacing = 1.15
                        p.paragraph_format.space_after = Pt(2)
                        p.paragraph_format.space_before = Pt(2)
                        
                        is_header = (r_idx == 0)
                        if is_header:
                            set_cell_background(cell, "F2F2F2")
                            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                            add_formatted_text(p, cols[c_idx], base_font_size=10, base_bold=True)
                        else:
                            # Align numbers or short codes center
                            if cols[c_idx].isdigit() or cols[c_idx] in ["Tersedia", "Tidak ada", "Terbatas", "Lolos", "Setuju", "Sangat Setuju"]:
                                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                            else:
                                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                            add_formatted_text(p, cols[c_idx], base_font_size=10)

                # Space after table
                p_after_tbl = doc.add_paragraph()
                p_after_tbl.paragraph_format.space_before = Pt(2)
                p_after_tbl.paragraph_format.space_after = Pt(6)
            continue

        # Check for Image
        if stripped.startswith("<p align=\"center\"><img"):
            m = re.search(r'src="([^"]+)"', line)
            src = m.group(1) if m else ""
            img_full_path = os.path.join(target_dir, src)
            if os.path.exists(img_full_path):
                p_img = doc.add_paragraph()
                p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p_img.paragraph_format.space_before = Pt(12)
                p_img.paragraph_format.space_after = Pt(4)
                run_img = p_img.add_run()
                with open(img_full_path, "rb") as f_img_stream:
                    run_img.add_picture(f_img_stream, width=Cm(14.0))

                # Look ahead for caption paragraph, skipping empty lines
                caption_text = lampiran_captions.get(src, "")
                next_idx = i + 1
                while next_idx < len(lines) and not lines[next_idx].strip():
                    next_idx += 1
                if next_idx < len(lines) and lines[next_idx].strip().startswith("**Lampiran"):
                    caption_text = lines[next_idx].strip()
                    i = next_idx + 1
                else:
                    i += 1
                if caption_text:
                    p_cap = doc.add_paragraph()
                    p_cap.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                    p_cap.paragraph_format.line_spacing = 1.5
                    p_cap.paragraph_format.space_after = Pt(12)
                    add_formatted_text(p_cap, caption_text, base_font_size=12)
            continue

        # Check for Headings
        if stripped.startswith("#"):
            heading_level = len(re.match(r'^#+', stripped).group(0))
            heading_text = stripped.lstrip('#').strip()
            
            p_head = doc.add_paragraph()
            p_head.paragraph_format.line_spacing = 1.5
            p_head.paragraph_format.space_before = Pt(12)
            p_head.paragraph_format.space_after = Pt(4)
            p_head.paragraph_format.keep_with_next = True

            if heading_text.upper() == "ABSTRAK":
                p_head.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p_head.paragraph_format.space_before = Pt(18)
                p_head.paragraph_format.space_after = Pt(12)
                add_formatted_text(p_head, "ABSTRAK", base_font_size=12, base_bold=True)
            elif heading_level == 2:
                p_head.alignment = WD_ALIGN_PARAGRAPH.LEFT
                add_formatted_text(p_head, heading_text, base_font_size=12, base_bold=True)
            elif heading_level == 3:
                p_head.alignment = WD_ALIGN_PARAGRAPH.LEFT
                add_formatted_text(p_head, heading_text, base_font_size=12, base_bold=True)
            else: # level 4+
                p_head.alignment = WD_ALIGN_PARAGRAPH.LEFT
                add_formatted_text(p_head, heading_text, base_font_size=12, base_bold=True)
            
            i += 1
            continue

        if not stripped:
            i += 1
            continue

        # Check for table title lines: **Tabel X. ...**
        if re.match(r'^\*\*Tabel\s+\d+\..*\*\*$', stripped):
            p_title = doc.add_paragraph()
            p_title.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p_title.paragraph_format.line_spacing = 1.5
            p_title.paragraph_format.space_before = Pt(10)
            p_title.paragraph_format.space_after = Pt(3)
            p_title.paragraph_format.keep_with_next = True
            add_formatted_text(p_title, stripped, base_font_size=12, base_bold=True)
            i += 1
            continue

        # Check for list items
        list_match = re.match(r'^(\d+\.|\-)\s+(.*)$', stripped)
        if list_match:
            bullet_sym = list_match.group(1)
            item_text = list_match.group(2)
            
            p_item = doc.add_paragraph()
            p_item.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p_item.paragraph_format.line_spacing = 1.5
            p_item.paragraph_format.space_after = Pt(3)
            p_item.paragraph_format.space_before = Pt(0)
            p_item.paragraph_format.left_indent = Cm(0.63)
            p_item.paragraph_format.first_line_indent = Cm(-0.63)

            prefix = f"{bullet_sym} " if bullet_sym.endswith('.') else "• "
            add_formatted_text(p_item, f"{prefix}{item_text}", base_font_size=12)
            i += 1
            continue

        # Check for sub-list items (indented list)
        sub_list_match = re.match(r'^\s+(\-|\d+\.)\s+(.*)$', line)
        if sub_list_match:
            bullet_sym = sub_list_match.group(1)
            item_text = sub_list_match.group(2)
            
            p_sub = doc.add_paragraph()
            p_sub.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p_sub.paragraph_format.line_spacing = 1.5
            p_sub.paragraph_format.space_after = Pt(2)
            p_sub.paragraph_format.space_before = Pt(0)
            p_sub.paragraph_format.left_indent = Cm(1.27)
            p_sub.paragraph_format.first_line_indent = Cm(-0.63)

            prefix = f"{bullet_sym} " if bullet_sym.endswith('.') else "– "
            add_formatted_text(p_sub, f"{prefix}{item_text}", base_font_size=12)
            i += 1
            continue

        # Regular Paragraph
        p_body = doc.add_paragraph()
        p_body.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p_body.paragraph_format.line_spacing = 1.5
        p_body.paragraph_format.space_after = Pt(3)
        p_body.paragraph_format.space_before = Pt(0)
        add_formatted_text(p_body, stripped, base_font_size=12)
        i += 1

    # Save to file
    doc.save(docx_path)
    file_size = os.path.getsize(docx_path)
    print(f"Sukses mengompilasi: '{docx_path}' ({file_size} bytes).")
    print(f"Total tabel dibuat: {len(doc.tables)} (1 cover + {data_table_count} data).")
    print(f"Total paragraf: {len(doc.paragraphs)}.")

if __name__ == "__main__":
    main()
